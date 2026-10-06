"""nozzle_report.py の出力 (report.json + fig_*.png) から PowerPoint を作る。`.venv-pptx` (python-pptx) で動かす。

usage: .venv-pptx/bin/python design/forge_design/report/build_pptx.py REPORT_DIR
規約の正本: procedures/nozzle-design-outputs.md
"""
import json
import sys
from datetime import date
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


def _wall_res_text(wr) -> str:
    """壁解像の正式値 (nozzle_report.wall_resolution = check_wall_resolution.py) を 1 行に。無ければ未評価と書く。"""
    if not wr:
        return "未評価 (report.json に wall_resolution が無い — nozzle_report を再実行)"
    parts = [f"VERDICT {wr.get('verdict')}"]
    if wr.get("over_area_pct") is not None:
        parts.append(f"y₁⁺ > {wr.get('target', 1):g} の面積 {wr['over_area_pct']:.1f} % (許容 {wr.get('over_area_allow_pct', float('nan')):g} %)")
    for name, w in (wr.get("per_wall") or {}).items():
        loc = f" (x/r_t = {w['x_max_rt']:.2f})" if "x_max_rt" in w else ""
        parts.append(f"{name}: 最大 {w['y1p_max']:.2f}{loc}・平均 {w['y1p_mean']:.2f}・p99 {w['y1p_p99']:.2f}")
    return "、".join(parts)

# 評価量の表の行 → --verdicts のキー (nozzle_report.QS_KEYS と同じ。判定するのは r/r_w = 0.1 の線)
QS_ROW_KEY = {"wave_pct": "wave_eta0.1", "overshoot_pct": "overshoot_eta0.1", "slope_pct": "slope_eta0.1",
              "range_pct": "range_eta0.1", "exit_core": "exit_core_M", "mdot": "mdot_ratio", "condensation": "condensation"}
QS_NONE = "VERDICT 未指定"


def _qs_cells(qs, key, tail=None):
    """評価量の表の VERDICT / 判定区間 / 備考 の 3 列。qs = report.json の quasisteady_verdicts (None = 未指定)。
    tail: 末尾 5 枚の幅の文字列 (従来の掲載値; VERDICT の有無によらず備考に残す)。"""
    it = ((qs or {}).get("items") or {}).get(key)
    note = [tail] if tail else []
    if not it:
        return [QS_NONE, "—", "; ".join(note) or "—"]
    if it.get("note"):
        note.append(it["note"])
    return [it["verdict"], it.get("window") or "—", "; ".join(note) or "—"]


def _qs_summary(qs) -> str:
    """判定ゲート表の 1 行: 量別の VERDICT と判定区間。未指定ならそう書く。"""
    items = (qs or {}).get("items") or {}
    if not items:
        return f"{QS_NONE} (nozzle_report --verdicts で check_quasisteady.py の判定を与える)。下の末尾 5 枚の幅は判定ではない"
    return "、".join(f"{k} {v['verdict']}" + (f" [{v['window']}]" if v.get("window") else "") for k, v in items.items())


NAVY = RGBColor(0x1B, 0x2B, 0x3A)
TEAL = RGBColor(0x0E, 0x7C, 0x86)
INK = RGBColor(0x22, 0x2B, 0x33)
MUTED = RGBColor(0x5E, 0x6B, 0x78)
PALE = RGBColor(0xEE, 0xF3, 0xF5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
WARN = RGBColor(0xB4, 0x5A, 0x09)
LATIN, EA = "Calibri", "Meiryo"
W, H = 13.333, 7.5


def _font(run, size, bold=False, color=INK):
    run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = color; run.font.name = LATIN
    rpr = run._r.get_or_add_rPr()
    ea = rpr.find(qn("a:ea"))
    if ea is None:
        ea = rpr.makeelement(qn("a:ea"), {}); rpr.append(ea)
    ea.set("typeface", EA)


def text(slide, x, y, w, h, lines, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    for i, ln in enumerate([lines] if isinstance(lines, str) else lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(4)
        segs = ln if isinstance(ln, list) else [(ln, {})]
        for seg, opt in segs:
            r = p.add_run(); r.text = seg
            _font(r, opt.get("size", size), opt.get("bold", bold), opt.get("color", color))
    return tb


def title(slide, t, sub=None):
    text(slide, 0.5, 0.3, W - 1.0, 0.7, t, size=28, bold=True, color=NAVY)
    if sub:
        text(slide, 0.5, 0.95, W - 1.0, 0.4, sub, size=12, color=MUTED)


def picture(slide, path, x, y, w, h):
    im = Image.open(path); ar = im.width / im.height
    if w / h > ar:
        ww, hh = h * ar, h
    else:
        ww, hh = w, w / ar
    slide.shapes.add_picture(str(path), Inches(x + (w - ww) / 2), Inches(y + (h - hh) / 2), Inches(ww), Inches(hh))


def table(slide, rows, x, y, w, col_w, header=True, size=11, row_h=0.32):
    nr, nc = len(rows), len(rows[0])
    shp = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
    tb = shp.table
    for j, cw in enumerate(col_w):
        tb.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = tb.cell(i, j); c.text = ""; p = c.text_frame.paragraphs[0]
            r = p.add_run(); r.text = str(val)
            hdr = header and i == 0
            _font(r, size, bold=hdr, color=WHITE if hdr else INK)
            c.fill.solid(); c.fill.fore_color.rgb = TEAL if hdr else (PALE if i % 2 == 0 else WHITE)
            c.margin_left = c.margin_right = Inches(0.06); c.margin_top = c.margin_bottom = Inches(0.03)
    return shp


def stat(slide, x, y, value, label, color=TEAL):
    box = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(2.9), Inches(1.35))
    box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0x24, 0x3B, 0x4E); box.line.fill.background()
    text(slide, x + 0.15, y + 0.1, 2.6, 0.7, value, size=28, bold=True, color=WHITE)
    text(slide, x + 0.15, y + 0.8, 2.6, 0.5, label, size=11, color=RGBColor(0xC9, 0xD6, 0xDF))


def main(report_dir):
    rd = Path(report_dir); rep = json.loads((rd / "report.json").read_text())
    cond, met, figs = rep["conditions"], rep["metrics"], rep["figures"]
    run_name, case = cond["run_name"], cond["case"]
    prs = Presentation(); prs.slide_width = Inches(W); prs.slide_height = Inches(H)
    blank = prs.slide_layouts[6]
    src = f"{case}/{run_name}/{cond['res']}"

    def new(t, sub=None, notes=None):
        s = prs.slides.add_slide(blank); title(s, t, sub)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    # 1 表紙
    s = prs.slides.add_slide(blank)
    bg = s.background.fill; bg.solid(); bg.fore_color.rgb = NAVY
    text(s, 0.7, 1.0, W - 1.4, 1.0, "ノズル設計 結果報告", size=40, bold=True, color=WHITE)
    text(s, 0.7, 1.9, W - 1.4, 0.6, f"{case} / {run_name}", size=18, color=RGBColor(0xC9, 0xD6, 0xDF))
    e01 = met["eta0.1"]; ec = met["exit_core"]
    qs = rep.get("quasisteady_verdicts")
    wv = ((qs or {}).get("items") or {}).get("wave_eta0.1")
    stat(s, 0.7, 3.2, f"{ec['M_mean']:.4f}", f"出口コアの M (設計 {met['Md']}; r/r_w 0.05〜0.7 の平均)")
    stat(s, 3.8, 3.2, f"{e01['wave_pct']:.3f} %",
         "試験部の Mach 波 (r/r_w=0.1、局所の山谷)" + (f" — {wv['verdict']}" if wv else f" — {QS_NONE}"))
    stat(s, 6.9, 3.2, f"{e01['overshoot_pct']:+.3f} %", f"オーバーシュート (M/M_d−1 の最大、x={e01['x_overshoot']:.0f})")
    cd = met.get("condensation")
    stat(s, 10.0, 3.2, f"{met['wall_shape']['exit_radius_m']:.4f} m" if not cd else f"x = {cd['onset_x_axis']:.1f}" if cd.get("onset_x_axis") else "—",
         "出口半径 (物理壁)" if not cd else "凝縮の始まり (軸、g > 1e-4)")
    text(s, 0.7, 6.6, W - 1.4, 0.5, f"作成 {date.today().isoformat()} · 出典 {src} · 生成 design/forge_design/report/nozzle_report.py",
         size=10, color=RGBColor(0x9F, 0xB2, 0xC0))

    # 2 解析領域
    s = new("解析領域と境界", "軸対称の上半分。色の線が境界パッチ (physID)", notes=f"図: {figs['domain']} ({src} の nozzle.h5)")
    picture(s, rd / figs["domain"], 0.5, 1.4, W - 1.0, 5.8)

    # 3 境界条件
    s = new("境界条件", "bcondConfig.yaml から転記")
    table(s, [["physID", "境界", "種類 (kind)", "値"]] + cond["bc"], 0.5, 1.5, W - 1.0, [1.0, 1.6, 2.4, W - 1.0 - 5.0], size=13, row_h=0.45)

    # 4 解析設定と物性
    s = new("解析設定と物性", "solverConfig.yaml・prepare_info.json から転記")
    table(s, [["項目", "設定"]] + cond["numerics"], 0.5, 1.4, 7.4, [1.6, 5.8], size=10.5, row_h=0.42)
    table(s, [["項目", "値"]] + cond["physics"], 8.2, 1.4, 4.6, [1.5, 3.1], size=10.5, row_h=0.42)

    # 5 ゲート
    s = new("判定ゲート", "合格したものだけでなく、不合格・未収束もそのまま載せる")
    gates = [["判定", "結果"]] + cond["gates"] + [
        ["準定常 VERDICT (量別、check_quasisteady.py)", _qs_summary(qs)],
        ["準定常 (末尾 5 枚の変動)", f"オーバーシュート {met['tail5_range']['overshoot_pct']:.4f} %、Mach 波 {met['tail5_range']['wave_pct']:.4f} %、出口コア M {met['tail5_range']['exit_core_M']:.5f}"],
        ["壁解像 y₁⁺ (check_wall_resolution.py、全 no-slip 壁)", _wall_res_text(met.get("wall_resolution"))]]
    table(s, gates, 0.5, 1.4, W - 1.0, [3.6, W - 1.0 - 3.6], size=11, row_h=0.6)

    # 6 評価量 (表) — 量別の準定常 VERDICT・判定区間・備考 (未達のまま進めた決定など) を同じ行に載せる
    t5 = met["tail5_range"]
    tails = {"wave_pct": f"末尾 {t5['n']} 枚の幅 {t5['wave_pct']:.4f} %pt", "overshoot_pct": f"末尾 {t5['n']} 枚の幅 {t5['overshoot_pct']:.4f} %pt",
             "exit_core": f"末尾 {t5['n']} 枚の幅 {t5['exit_core_M']:.5f}"}
    s = new("評価量 — 試験部の一様性", "軸 M は形状の生成器。見るのは Mach 波・オーバーシュート・傾き (r/r_w=0.1 が軸ノードの癖を避けた線)。"
            "VERDICT は r/r_w=0.1 の量の準定常判定",
            notes=f"試験部 = x ∈ [{met['eta0.1']['test_window'][0]:.1f}, {met['eta0.1']['test_window'][1]:.1f}] r_t。"
                  f"VERDICT の出典: {(qs or {}).get('source') or '未指定'}")
    rows = [["量", "r = 0 (軸)", "r/r_w = 0.1", "VERDICT", "判定区間", "備考"]]
    for k, lab in (("wave_pct", "Mach 波 (100(M/M_d−1) の 10 r_t 平滑からの残差の最大) [%]"), ("overshoot_pct", "オーバーシュート M/M_d−1 の最大 [%]"),
                   ("slope_pct", "試験部の傾き (1 次近似の両端差) [%]"), ("range_pct", "試験部の最大−最小 [%]")):
        rows.append([lab, f"{met['eta0.0'][k]:+.4f}", f"{met['eta0.1'][k]:+.4f}"] + _qs_cells(qs, QS_ROW_KEY[k], tails.get(k)))
    rows.append(["出口コア M (平均 / 最小〜最大)", f"{ec['M_mean']:.5f}", f"{ec['M_min']:.5f}〜{ec['M_max']:.5f}"]
                + _qs_cells(qs, QS_ROW_KEY["exit_core"], tails["exit_core"]))
    if "mdot_ratio_vs_euler" in met:
        v = met["mdot_ratio_vs_euler"]
        rows.append(["流量比 ṁ_NS / ṁ_Euler", f"{v:.6f}" if isinstance(v, float) else str(v), ""] + _qs_cells(qs, QS_ROW_KEY["mdot"]))
    if cd:
        rows.append(["凝縮: 始まり (軸) / 出口 g (軸・コア平均)", f"x = {cd['onset_x_axis']}" if cd.get("onset_x_axis") else "—",
                     f"{cd['exit_g_axis']:.4f} / {cd['exit_g_core_mean']:.4f}"] + _qs_cells(qs, QS_ROW_KEY["condensation"]))
    # 表の行に対応しない量 (δ_E/δ_C・凝縮の個別量など) も落とさず別行で載せる
    for k in ((qs or {}).get("items") or {}):
        if k not in QS_ROW_KEY.values():
            rows.append([k, "", ""] + _qs_cells(qs, k))
    table(s, rows, 0.5, 1.5, W - 1.0, [3.3, 1.1, 1.5, 1.6, 1.4, W - 1.0 - 8.9], size=10, row_h=0.4)

    # 6′ 評価量 (図)
    s = new("評価量 — 試験部の M/M_d − 1", "r = 0 (実線) と r/r_w = 0.1 (破線)、黒は Euler 参照 (設計壁)",
            notes=f"図: {figs['axis_dev']} ({src})")
    picture(s, rd / figs["axis_dev"], 0.5, 1.4, W - 1.0, 5.9)

    # 7.. コンタ
    names = {"fig_contour_mach.png": ("コンタ: マッハ数", "上: M (全域)、下: M/M_d − 1 (±0.5 %)"),
             "fig_contour_gradients.png": ("コンタ: 圧力勾配とシュリーレン", "上: (r_t/p) ∂p/∂x、下: log10(|∇ρ| r_t/ρ)"),
             "fig_contour_PT.png": ("コンタ: 静圧と静温", "0.2〜99.8 % で頭打ち"),
             "fig_contour_condensation.png": ("コンタ: 凝縮", "上から 液滴の質量分率 g、過冷却度 T_sat − T、過飽和度 S")}
    for f in figs["contours"]:
        t, sub = names.get(f, (f, None))
        s = new(t, sub, notes=f"図: {f} ({src})。カラーマップは全量 turbo")
        picture(s, rd / f, 0.5, 1.4, W - 1.0, 5.9)

    for key, t, sub in (("axis", "軸に沿った分布", "M・静圧・静温・密度・動圧 0.5ρv²・0.5v²・全圧 (凝縮ありは g・過冷却度・S)。実線 r=0、破線 r/r_w=0.1"),
                        ("exit", "出口断面の分布", "同じ量を半径方向に。流れ角 atan(v/u) を追加"),
                        ("wall", "壁面の分布", "壁圧・壁温・C_f・y₁⁺ 近似 (図示のみ。y₁⁺ の判定値は判定ゲート表 = check_wall_resolution.py)"),
                        ("wall_shape", "壁の形と微分", "左: r・r′・r″ (点線 = 設計壁)、右: 設計壁との差 (境界層の補正分)")):
        s = new(t, sub, notes=f"図: {figs[key]} ({src})")
        picture(s, rd / figs[key], 0.5, 1.4, W - 1.0, 5.9)

    out = rd / f"{run_name}_report.pptx"
    prs.save(out)
    print(f"[build_pptx] {out} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main(sys.argv[1])
